#!/usr/bin/env python3
"""Export the RADE overview as a genuinely editable one-slide PowerPoint.

All diagram text, lines, arrows, blocks, cube outlines, and labels are native
PowerPoint objects.  Only the six measured radar heatmap surfaces are raster
picture objects, each separately selectable in PowerPoint.

Install ``python-pptx`` before running this exporter.  The generated deck is
written beside the existing PNG/SVG/PDF exports and does not modify them.
"""

from __future__ import annotations

import argparse
from io import BytesIO
import os
from pathlib import Path


os.environ.setdefault("MPLCONFIGDIR", "/tmp/mvrss_rade_pptx_matplotlib")

from matplotlib.backends.backend_agg import FigureCanvasAgg
from matplotlib.figure import Figure
from matplotlib.transforms import Affine2D
import numpy as np

try:
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.dml import MSO_LINE_DASH_STYLE
    from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
    from pptx.enum.text import MSO_VERTICAL_ANCHOR, PP_ALIGN
    from pptx.oxml.ns import qn
    from pptx.oxml.xmlchemy import OxmlElement
    from pptx.util import Inches, Pt
except ModuleNotFoundError as exc:  # pragma: no cover - environment guidance
    raise SystemExit(
        "python-pptx is required. Install it with: python -m pip install python-pptx"
    ) from exc

from .architecture import DEFAULT_RADAR_ROOT, ROOT
from .overall_process import load_projection_maps


OUT_DIR = ROOT / "figures" / "multiview_rade_overview"
OUT_PATH = OUT_DIR / "rade_to_rad_rae_overall_process_editable.pptx"

SLIDE_WIDTH_IN = 16.9
SLIDE_HEIGHT_IN = 7.244
CANVAS_WIDTH = 216.0
CANVAS_HEIGHT = 92.0
X_MARGIN_IN = 0.08
Y_MARGIN_IN = 0.08
X_SCALE_IN = 0.0775
Y_SCALE_IN = 0.0770

WHITE = "FFFFFF"
INK = "263746"
NAVY = "17324D"
MID = "687783"
LIGHT = "D9E2E8"
BLUE = "2474B5"
PALE_BLUE = "E9F3FA"
PURPLE = "6F5AAF"
PALE_PURPLE = "F0EDF8"
GREEN = "378E74"
PALE_GREEN = "EAF5F0"
CUBE_EDGE = "33495C"
CUBE_LIGHT = "E7E9EC"
CUBE_MID = "D7DADE"
CUBE_DARK = "BFC4CA"
FONT = "Georgia"
RA_CMAP = "viridis"


def rgb(value: str) -> RGBColor:
    return RGBColor.from_string(value.lstrip("#"))


def sx(value: float) -> int:
    return Inches(X_MARGIN_IN + X_SCALE_IN * value)


def sy(value: float) -> int:
    """Convert a bottom-up source y-coordinate to a slide y-coordinate."""
    return Inches(Y_MARGIN_IN + Y_SCALE_IN * (CANVAS_HEIGHT - value))


def box_emu(x: float, y: float, width: float, height: float) -> tuple[int, int, int, int]:
    return (
        sx(x),
        sy(y + height),
        Inches(X_SCALE_IN * width),
        Inches(Y_SCALE_IN * height),
    )


def set_shape_fill(shape, colour: str | None) -> None:
    if colour is None:
        shape.fill.background()
        return
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(colour)


def set_shape_line(
    shape,
    colour: str | None,
    *,
    width_pt: float = 1.0,
    dash: MSO_LINE_DASH_STYLE | None = None,
) -> None:
    if colour is None:
        shape.line.fill.background()
        return
    shape.line.fill.solid()
    shape.line.fill.fore_color.rgb = rgb(colour)
    shape.line.width = Pt(width_pt)
    if dash is not None:
        shape.line.dash_style = dash


def set_end_arrow(connector) -> None:
    """Add a native DrawingML triangular arrowhead at connector end."""
    line_xml = connector._element.spPr.ln
    old = line_xml.find(qn("a:tailEnd"))
    if old is not None:
        line_xml.remove(old)
    tail = OxmlElement("a:tailEnd")
    tail.set("type", "triangle")
    tail.set("w", "sm")
    tail.set("len", "sm")
    line_xml.append(tail)


def add_line(
    slide,
    start: tuple[float, float],
    end: tuple[float, float],
    *,
    colour: str = INK,
    width_pt: float = 1.0,
    dash: bool = False,
    arrowhead: bool = False,
    name: str,
):
    connector = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT,
        sx(start[0]),
        sy(start[1]),
        sx(end[0]),
        sy(end[1]),
    )
    connector.name = name
    set_shape_line(
        connector,
        colour,
        width_pt=width_pt,
        dash=MSO_LINE_DASH_STYLE.DASH if dash else MSO_LINE_DASH_STYLE.SOLID,
    )
    if arrowhead:
        set_end_arrow(connector)
    return connector


def add_polygon(
    slide,
    points: list[tuple[float, float]] | tuple[tuple[float, float], ...],
    *,
    fill: str | None,
    line: str | None,
    width_pt: float = 1.0,
    dash: bool = False,
    name: str,
):
    vertices = [(sx(x), sy(y)) for x, y in points]
    builder = slide.shapes.build_freeform(*vertices[0])
    builder.add_line_segments(vertices[1:], close=True)
    shape = builder.convert_to_shape()
    shape.name = name
    set_shape_fill(shape, fill)
    set_shape_line(
        shape,
        line,
        width_pt=width_pt,
        dash=MSO_LINE_DASH_STYLE.DASH if dash else MSO_LINE_DASH_STYLE.SOLID,
    )
    return shape


def style_text_frame(
    shape,
    text: str,
    *,
    size_pt: float,
    colour: str,
    bold: bool = False,
    italic: bool = False,
    align: PP_ALIGN = PP_ALIGN.CENTER,
) -> None:
    frame = shape.text_frame
    frame.clear()
    frame.margin_left = 0
    frame.margin_right = 0
    frame.margin_top = 0
    frame.margin_bottom = 0
    frame.vertical_anchor = MSO_VERTICAL_ANCHOR.MIDDLE
    frame.word_wrap = True
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    run = paragraph.add_run()
    run.text = text
    run.font.name = FONT
    run.font.size = Pt(size_pt)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = rgb(colour)


def add_text(
    slide,
    x: float,
    y: float,
    width: float,
    height: float,
    text: str,
    *,
    size_pt: float,
    colour: str = INK,
    bold: bool = False,
    italic: bool = False,
    align: PP_ALIGN = PP_ALIGN.CENTER,
    rotation: float = 0.0,
    name: str,
):
    shape = slide.shapes.add_textbox(*box_emu(x, y, width, height))
    shape.name = name
    shape.rotation = rotation
    style_text_frame(
        shape,
        text,
        size_pt=size_pt,
        colour=colour,
        bold=bold,
        italic=italic,
        align=align,
    )
    return shape


def add_text_center(
    slide,
    center_x: float,
    center_y: float,
    width: float,
    height: float,
    text: str,
    **kwargs,
):
    return add_text(
        slide,
        center_x - width / 2,
        center_y - height / 2,
        width,
        height,
        text,
        **kwargs,
    )


def add_rounded_text_box(
    slide,
    x: float,
    y: float,
    width: float,
    height: float,
    text: str,
    *,
    fill: str,
    line: str,
    text_colour: str,
    size_pt: float,
    bold: bool,
    name: str,
):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, *box_emu(x, y, width, height))
    shape.name = name
    set_shape_fill(shape, fill)
    set_shape_line(shape, line, width_pt=0.75)
    style_text_frame(
        shape,
        text,
        size_pt=size_pt,
        colour=text_colour,
        bold=bold,
    )
    return shape


def render_face_png(
    image: np.ndarray,
    *,
    bbox_width: float,
    bbox_height: float,
    corners: tuple[tuple[float, float], tuple[float, float], tuple[float, float], tuple[float, float]],
) -> bytes:
    """Render one heatmap into a transparent face-aligned PNG."""
    fig = Figure(
        figsize=(max(0.15, X_SCALE_IN * bbox_width), max(0.15, Y_SCALE_IN * bbox_height)),
        dpi=360,
        facecolor=(1, 1, 1, 0),
    )
    FigureCanvasAgg(fig)
    axis = fig.add_axes((0, 0, 1, 1), frameon=False)
    axis.set_xlim(0.0, bbox_width)
    axis.set_ylim(0.0, bbox_height)
    axis.axis("off")

    p00 = np.asarray(corners[0], dtype=float)
    p10 = np.asarray(corners[1], dtype=float)
    p01 = np.asarray(corners[3], dtype=float)
    basis_u = p10 - p00
    basis_v = p01 - p00
    transform = Affine2D.from_values(
        basis_u[0],
        basis_u[1],
        basis_v[0],
        basis_v[1],
        p00[0],
        p00[1],
    ) + axis.transData
    axis.imshow(
        image,
        extent=(0.0, 1.0, 0.0, 1.0),
        origin="lower",
        aspect="auto",
        cmap=RA_CMAP,
        vmin=0.0,
        vmax=1.0,
        interpolation="bilinear",
        transform=transform,
    )
    buffer = BytesIO()
    fig.savefig(buffer, format="png", transparent=True, dpi=360)
    return buffer.getvalue()


def add_picture_bytes(
    slide,
    data: bytes,
    x: float,
    y: float,
    width: float,
    height: float,
    *,
    name: str,
):
    picture = slide.shapes.add_picture(BytesIO(data), *box_emu(x, y, width, height))
    picture.name = name
    return picture


def cube_vertices(
    origin: np.ndarray,
    u: np.ndarray,
    v: np.ndarray,
    w: np.ndarray,
) -> list[np.ndarray]:
    return [origin + a * u + b * v + c * w for c in (0, 1) for b in (0, 1) for a in (0, 1)]


def cube_edges(vertices: list[np.ndarray]) -> list[tuple[np.ndarray, np.ndarray]]:
    edges: list[tuple[np.ndarray, np.ndarray]] = []
    for index in range(8):
        flags = (index & 1, (index >> 1) & 1, (index >> 2) & 1)
        for bit, flag in zip((1, 2, 4), flags):
            if flag:
                continue
            edges.append((vertices[index], vertices[index | bit]))
    if len(edges) != 12:
        raise AssertionError(f"A cube must have 12 edges, got {len(edges)}")
    return edges


def draw_four_axis_glyph(slide) -> None:
    origin = np.asarray((8.5, 29.5), dtype=float)
    azimuth = np.asarray((12.5, 0.0), dtype=float)
    range_axis = np.asarray((4.2, 4.8), dtype=float)
    elevation = np.asarray((0.0, 12.5), dtype=float)
    doppler = np.asarray((-4.2, 4.8), dtype=float)
    near = cube_vertices(origin, azimuth, range_axis, elevation)
    far = [point + doppler for point in near]

    face_specs = (
        ((0, 1, 5, 4), CUBE_LIGHT),
        ((4, 5, 7, 6), CUBE_MID),
        ((1, 3, 7, 5), CUBE_DARK),
    )
    for cube_name, vertices in (("Far", far), ("Near", near)):
        for face_index, (indices, colour) in enumerate(face_specs, start=1):
            add_polygon(
                slide,
                [tuple(vertices[index]) for index in indices],
                fill=colour,
                line=None,
                name=f"4D_{cube_name}Cube_Face_{face_index}",
            )

    for index, (near_point, far_point) in enumerate(zip(near, far), start=1):
        add_line(
            slide,
            tuple(near_point),
            tuple(far_point),
            colour=CUBE_EDGE,
            width_pt=0.8,
            dash=True,
            name=f"4D_Doppler_Link_{index:02d}",
        )
    for cube_name, vertices in (("Far", far), ("Near", near)):
        for edge_index, (start, end) in enumerate(cube_edges(vertices), start=1):
            add_line(
                slide,
                tuple(start),
                tuple(end),
                colour=CUBE_EDGE,
                width_pt=0.9,
                name=f"4D_{cube_name}Cube_Edge_{edge_index:02d}",
            )

    add_line(slide, (8.5, 27.8), (21.0, 27.8), colour="000000", width_pt=0.9, arrowhead=True, name="Axis_Azimuth")
    add_text_center(slide, 14.75, 25.9, 10.0, 2.0, "Azimuth", size_pt=7.2, colour="000000", name="Label_Azimuth")
    add_line(slide, (7.8, 29.3), (2.55, 35.3), colour="000000", width_pt=0.9, arrowhead=True, name="Axis_Doppler")
    add_text_center(slide, 4.1, 30.8, 8.0, 2.0, "Doppler", size_pt=7.2, colour="000000", name="Label_Doppler")
    add_line(slide, (22.1, 28.3), (27.8, 34.8), colour="000000", width_pt=0.9, arrowhead=True, name="Axis_Range")
    add_text_center(slide, 25.6, 29.7, 8.0, 2.0, "Range", size_pt=7.2, colour="000000", name="Label_Range")
    add_line(slide, (28.6, 34.8), (28.6, 46.7), colour="000000", width_pt=0.9, arrowhead=True, name="Axis_Elevation")
    add_text_center(slide, 30.6, 40.8, 12.0, 2.0, "Elevation", size_pt=7.2, colour="000000", rotation=270.0, name="Label_Elevation")
    add_text_center(slide, 13.5, 56.0, 24.0, 3.0, "4-D RADE tensor", size_pt=10.2, colour=NAVY, bold=True, name="Title_4D_RADE")


def draw_tensor_cube(
    slide,
    x: float,
    y: float,
    maps: dict[str, np.ndarray],
    *,
    cube_name: str,
    side_view: str,
    top_view: str,
    shape_text: str,
    shared_front_png: bytes,
) -> None:
    width, height, depth, rise = 18.5, 18.0, 8.0, 4.8
    top_local = ((0.0, 0.0), (width, 0.0), (width + depth, rise), (depth, rise))
    side_local = ((0.0, 0.0), (depth, rise), (depth, height + rise), (0.0, height))
    top_png = render_face_png(
        maps[top_view].T,
        bbox_width=width + depth,
        bbox_height=rise,
        corners=top_local,
    )
    side_png = render_face_png(
        maps[side_view],
        bbox_width=depth,
        bbox_height=height + rise,
        corners=side_local,
    )
    add_picture_bytes(slide, top_png, x, y + height, width + depth, rise, name=f"{cube_name}_{top_view}_Surface")
    add_picture_bytes(slide, side_png, x + width, y, depth, height + rise, name=f"{cube_name}_{side_view}_Surface")
    add_picture_bytes(slide, shared_front_png, x, y, width, height, name=f"{cube_name}_RA_Surface")

    front = [(x, y), (x + width, y), (x + width, y + height), (x, y + height)]
    top = [(x, y + height), (x + width, y + height), (x + width + depth, y + height + rise), (x + depth, y + height + rise)]
    side = [(x + width, y), (x + width + depth, y + rise), (x + width + depth, y + height + rise), (x + width, y + height)]
    add_polygon(slide, top, fill=None, line=CUBE_EDGE, width_pt=1.0, name=f"{cube_name}_Top_Outline")
    add_polygon(slide, side, fill=None, line=CUBE_EDGE, width_pt=1.0, name=f"{cube_name}_Side_Outline")
    add_polygon(slide, front, fill=None, line=CUBE_EDGE, width_pt=1.0, name=f"{cube_name}_Front_Outline")

    add_rounded_text_box(slide, x + 7.1, y + 0.65, 4.3, 1.7, "RA view", fill=WHITE, line=CUBE_EDGE, text_colour=NAVY, size_pt=6.4, bold=True, name=f"{cube_name}_RA_Label")
    add_rounded_text_box(slide, x + 10.0, y + height + 1.45, 5.0, 1.7, f"{top_view} view", fill=WHITE, line=CUBE_EDGE, text_colour=NAVY, size_pt=6.4, bold=True, name=f"{cube_name}_{top_view}_Label")
    side_label = add_rounded_text_box(slide, x + width + 2.8, y + 8.0, 5.0, 1.7, f"{side_view} view", fill=WHITE, line=CUBE_EDGE, text_colour=NAVY, size_pt=6.4, bold=True, name=f"{cube_name}_{side_view}_Label")
    side_label.rotation = 270.0
    add_text_center(slide, x + width / 2 + depth / 2, y + height + rise + 2.5, 18.0, 2.5, f"{cube_name} tensor", size_pt=10.0, colour=BLUE, bold=True, name=f"{cube_name}_Title")
    add_text_center(slide, x + width / 2, y - 2.3, 12.0, 2.0, shape_text, size_pt=7.2, colour=INK, bold=True, name=f"{cube_name}_Shape")


def draw_encoder(slide, x: float, y: float, title: str) -> None:
    width, height = 22.0, 18.0
    points = [(x, y), (x + width, y + 3.2), (x + width, y + height - 3.2), (x, y + height)]
    add_polygon(slide, points, fill=PALE_BLUE, line=BLUE, width_pt=1.15, name=f"{title}_Encoder_Block")
    for index, offset in enumerate((5.5, 11.0, 16.5), start=1):
        lower = y + offset * 3.2 / width
        upper = y + height - offset * 3.2 / width
        add_line(slide, (x + offset, lower), (x + offset, upper), colour=BLUE, width_pt=0.45, name=f"{title}_Encoder_Stage_{index}")
    add_text_center(slide, x + 10.4, y + 10.1, 17.0, 2.5, f"{title} encoder", size_pt=9.0, colour=BLUE, bold=True, name=f"{title}_Encoder_Title")
    add_text_center(slide, x + 10.4, y + 6.7, 12.0, 2.0, "Swin–FPN", size_pt=7.0, colour=INK, name=f"{title}_Encoder_Subtitle")


def draw_fusion(slide) -> None:
    x, y = 113.0, 32.0
    for index, (dx, dy, colour) in enumerate(((1.2, 1.5, PALE_BLUE), (0.6, 0.8, PALE_BLUE), (0.0, 0.0, PALE_PURPLE)), start=1):
        add_polygon(
            slide,
            [(x + dx, y + dy), (x + dx + 10.5, y + dy), (x + dx + 12.9, y + dy + 2.4), (x + dx + 12.9, y + dy + 24.4), (x + dx + 2.4, y + dy + 24.4), (x + dx, y + dy + 22.0)],
            fill=colour,
            line=PURPLE if index == 3 else BLUE,
            width_pt=0.85,
            name=f"Fusion_Latent_Slab_{index}",
        )
    add_text_center(slide, 118.25, 43.2, 9.5, 5.0, "encoded\nconcat + fusion", size_pt=7.6, colour=PURPLE, bold=True, name="Fusion_Label")
    add_text_center(slide, 118.25, 29.6, 10.0, 2.0, "shared latent", size_pt=6.6, colour=PURPLE, bold=True, name="Fusion_Subtitle")


def draw_decoder(slide) -> None:
    x, y, width, height = 134.0, 30.5, 22.0, 25.0
    points = [(x, y + 4.5), (x + width, y), (x + width, y + height), (x, y + height - 4.5)]
    add_polygon(slide, points, fill=PALE_PURPLE, line=PURPLE, width_pt=1.2, name="CenterPoint_Decoder_Block")
    for index, offset in enumerate((5.5, 11.0, 16.5), start=1):
        lower = y + 4.5 * (1.0 - offset / width)
        upper = y + height - 4.5 * (1.0 - offset / width)
        add_line(slide, (x + offset, lower), (x + offset, upper), colour=PURPLE, width_pt=0.45, name=f"Decoder_Stage_{index}")
    add_text_center(slide, x + width / 2, y + 15.3, 16.0, 2.5, "CenterPoint", size_pt=9.8, colour=PURPLE, bold=True, name="Decoder_Title")
    add_text_center(slide, x + width / 2, y + 11.5, 12.0, 2.3, "decoder", size_pt=9.0, colour=PURPLE, bold=True, name="Decoder_Subtitle")
    add_text_center(slide, x + width / 2, y + 7.5, 18.0, 2.0, "heatmap + box regression", size_pt=6.4, colour=INK, name="Decoder_Detail")


def draw_output(slide) -> None:
    x, y, width, height = 176.0, 28.5, 21.0, 29.0
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, *box_emu(x, y, width, height))
    card.name = "Detection_Output_Card"
    set_shape_fill(card, PALE_GREEN)
    set_shape_line(card, GREEN, width_pt=1.1)
    add_text_center(slide, x + width / 2, y + height - 3.2, 18.0, 2.5, "Detection output", size_pt=9.2, colour=GREEN, bold=True, name="Output_Title")
    add_text_center(slide, x + width / 2, y + height - 6.1, 12.0, 2.0, "placeholder", size_pt=7.0, colour=MID, bold=True, name="Output_Placeholder")
    plane = [(181.2, 37.3), (190.0, 40.7), (194.5, 37.5), (185.7, 34.2)]
    add_polygon(slide, plane, fill=WHITE, line=GREEN, width_pt=0.8, dash=True, name="Output_BEV_Plane")
    # Small editable wireframe 3-D box.
    lower = [(186.5, 37.2), (188.0, 37.8), (189.4, 37.1), (187.9, 36.5)]
    upper = [(x0, y0 + 1.6) for x0, y0 in lower]
    for index, ring in enumerate((lower, upper), start=1):
        add_polygon(slide, ring, fill=None, line=GREEN, width_pt=0.8, name=f"Output_Box_Ring_{index}")
    for index, (p0, p1) in enumerate(zip(lower, upper), start=1):
        add_line(slide, p0, p1, colour=GREEN, width_pt=0.8, name=f"Output_Box_Vertical_{index}")
    add_text_center(slide, x + width / 2, y + 2.0, 18.0, 2.0, "final visualization to be revised", size_pt=6.0, colour=MID, italic=True, name="Output_Note")


def add_stage_heading(slide, x: float, text: str, colour: str, name: str) -> None:
    add_text_center(slide, x, 76.7, 28.0, 2.2, text, size_pt=8.5, colour=colour, bold=True, name=name)


def build_presentation(rad_maps, rae_maps, source_note: str) -> Presentation:
    presentation = Presentation()
    presentation.slide_width = Inches(SLIDE_WIDTH_IN)
    presentation.slide_height = Inches(SLIDE_HEIGHT_IN)
    presentation.core_properties.title = "Overall RADE dual-view processing and detection"
    presentation.core_properties.subject = "Editable Model7 radar architecture overview"
    presentation.core_properties.comments = "Native PowerPoint geometry; measured radar surfaces are separate raster picture objects."
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = rgb(WHITE)

    add_text(slide, 4.0, 86.3, 115.0, 3.4, "Overall RADE dual-view processing and detection", size_pt=17.5, colour=NAVY, bold=True, align=PP_ALIGN.LEFT, name="Header_Title")
    add_text(slide, 4.1, 82.9, 152.0, 2.6, "Four-axis radar tensor  →  three-view RAD and RAE cubes  →  independent encoders  →  encoded-feature fusion  →  shared detector", size_pt=8.6, colour=MID, align=PP_ALIGN.LEFT, name="Header_Subtitle")
    add_rounded_text_box(slide, 197.0, 85.7, 15.0, 3.0, "CURRENT MODEL7", fill=PALE_BLUE, line=BLUE, text_colour=BLUE, size_pt=6.5, bold=True, name="Current_Model_Badge")
    add_line(slide, (4.0, 80.5), (212.0, 80.5), colour=LIGHT, width_pt=0.9, name="Header_Rule")

    add_stage_heading(slide, 14.0, "4-D INPUT", NAVY, "Stage_4D_Input")
    add_stage_heading(slide, 59.0, "MULTI-VIEW TENSORS", BLUE, "Stage_Multiview_Tensors")
    add_stage_heading(slide, 89.0, "DUAL ENCODERS", BLUE, "Stage_Dual_Encoders")
    add_stage_heading(slide, 144.0, "FUSION + DECODING", PURPLE, "Stage_Fusion_Decoding")
    add_stage_heading(slide, 187.0, "OUTPUT", GREEN, "Stage_Output")

    draw_four_axis_glyph(slide)

    add_line(slide, (29.5, 58.5), (45.5, 55.5), colour=BLUE, width_pt=1.2, arrowhead=True, name="Arrow_Mean_E")
    add_rounded_text_box(slide, 33.5, 61.2, 8.0, 1.9, "mean over E", fill=WHITE, line=BLUE, text_colour=BLUE, size_pt=6.3, bold=True, name="Label_Mean_E")
    add_line(slide, (29.5, 34.0), (45.5, 25.5), colour=BLUE, width_pt=1.2, arrowhead=True, name="Arrow_Mean_D")
    add_rounded_text_box(slide, 33.0, 35.7, 8.0, 1.9, "mean over D", fill=WHITE, line=BLUE, text_colour=BLUE, size_pt=6.3, bold=True, name="Label_Mean_D")

    front_png = render_face_png(
        rad_maps["RA"],
        bbox_width=18.5,
        bbox_height=18.0,
        corners=((0.0, 0.0), (18.5, 0.0), (18.5, 18.0), (0.0, 18.0)),
    )
    draw_tensor_cube(slide, 46.0, 46.5, rad_maps, cube_name="RAD", side_view="RD", top_view="AD", shape_text="[R, A, D]", shared_front_png=front_png)
    draw_tensor_cube(slide, 46.0, 16.5, rae_maps, cube_name="RAE", side_view="RE", top_view="AE", shape_text="[R, A, E]", shared_front_png=front_png)

    add_line(slide, (72.8, 55.5), (77.5, 55.5), colour=BLUE, width_pt=1.2, arrowhead=True, name="Arrow_RAD_to_Encoder")
    add_line(slide, (72.8, 25.5), (77.5, 25.5), colour=BLUE, width_pt=1.2, arrowhead=True, name="Arrow_RAE_to_Encoder")
    draw_encoder(slide, 78.0, 46.5, "RAD")
    draw_encoder(slide, 78.0, 16.5, "RAE")
    add_text_center(slide, 89.0, 43.0, 16.0, 2.0, "independent weights", size_pt=6.5, colour=MID, italic=True, name="Independent_Weights")

    add_line(slide, (100.3, 55.5), (112.5, 49.5), colour=BLUE, width_pt=1.2, arrowhead=True, name="Arrow_RAD_Encoder_to_Fusion")
    add_line(slide, (100.3, 25.5), (112.5, 39.5), colour=BLUE, width_pt=1.2, arrowhead=True, name="Arrow_RAE_Encoder_to_Fusion")
    draw_fusion(slide)
    add_line(slide, (126.5, 43.0), (133.5, 43.0), colour=PURPLE, width_pt=1.3, arrowhead=True, name="Arrow_Fusion_to_Decoder")
    draw_decoder(slide)
    add_line(slide, (156.5, 43.0), (175.5, 43.0), colour=BLUE, width_pt=1.2, arrowhead=True, name="Arrow_Decode")
    add_rounded_text_box(slide, 163.3, 45.0, 5.4, 1.9, "decode", fill=WHITE, line=BLUE, text_colour=BLUE, size_pt=6.3, bold=True, name="Decode_Label")
    draw_output(slide)

    add_line(slide, (4.0, 9.5), (212.0, 9.5), colour=LIGHT, width_pt=0.8, name="Footer_Rule")
    add_text(slide, 4.0, 5.6, 205.0, 2.0, f"Real paired data: {source_note}. Every visible cube surface uses the same colormap and display-normalization rule.", size_pt=6.4, colour=INK, align=PP_ALIGN.LEFT, name="Footer_Data_Provenance")
    add_text(slide, 4.0, 2.8, 205.0, 2.0, "RAD surfaces: RA front · RD side · AD top. RAE surfaces: RA front · RE side · AE top. Four-axis glyph adapted from 4DR P2T Fig. 2; artwork is original.", size_pt=6.1, colour=MID, italic=True, align=PP_ALIGN.LEFT, name="Footer_Surface_Key")

    return presentation


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output", type=Path, default=OUT_PATH)
    parser.add_argument("--radar-root", type=Path, default=DEFAULT_RADAR_ROOT)
    parser.add_argument("--sample-seed", type=int, default=42)
    parser.add_argument("--sequence", type=str, default=None)
    parser.add_argument("--frame", type=str, default=None)
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    rad_maps, rae_maps, source_note = load_projection_maps(
        args.radar_root,
        seed=args.sample_seed,
        sequence=args.sequence,
        frame=args.frame,
    )
    if not np.array_equal(rad_maps["RA"], rae_maps["RA"]):
        raise AssertionError("The two PowerPoint RA surfaces must share identical source pixels")
    presentation = build_presentation(rad_maps, rae_maps, source_note)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    presentation.save(args.output)
    print(source_note)
    print(args.output.resolve())


if __name__ == "__main__":
    main()
